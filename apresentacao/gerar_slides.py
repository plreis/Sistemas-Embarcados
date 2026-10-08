#!/usr/bin/env python3
"""Gera a apresentacao e o roteiro a partir dos arquivos da entrega.
Dependencias: python-pptx==1.0.2 e Pillow. PDF: exportar o PPTX no LibreOffice.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = Path(__file__).resolve().parent
REPO = BASE.parent
FOTOS = REPO / 'fotos'
FONT = 'DejaVu Sans'
MONO = 'DejaVu Sans Mono'
NAVY='142638'; INK='172A3A'; GOLD='F6C644'; GRAY='586977'; WHITE='FFFFFF'
BG='F7F9FB'; LIGHT='E8EEF3'; TEAL='147D80'; DARKCARD='233B4E'; PALE='E7F3F1'
P = Presentation(); P.slide_width=Inches(13.333); P.slide_height=Inches(7.5)
P.core_properties.title='Relógio e alarme com FreeRTOS'
P.core_properties.subject='Exame de Suficiência — Sistemas Embarcados 2026.2'
P.core_properties.author='Pedro Lucas dos Reis Silva'
P.core_properties.keywords='ATmega2560, FreeRTOS, filas, Timer1, suficiência'
SLIDES=[]

def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,color,line=None,radius=False):
 z=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 if radius:z.adjustments[0]=0.10
 z.fill.solid();z.fill.fore_color.rgb=rgb(color)
 if line:z.line.color.rgb=rgb(line);z.line.width=Pt(1)
 else:z.line.fill.background()
 return z

def txt(s,x,y,w,h,text,size=22,color=INK,bold=False,font=FONT,align=None):
 z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=z.text_frame
 tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
  p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
  p.space_after=Pt(2 if font==MONO else 8);p.space_before=Pt(0)
  if align is not None:p.alignment=align
 return z

def arrow(s,x,y,w=0.45,h=0.2,color=TEAL,down=False):
 a=s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW if down else MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
 a.fill.solid();a.fill.fore_color.rgb=rgb(color);a.line.fill.background();return a

def photo(s,name,x,y,w,h):
 from PIL import Image
 path=FOTOS/name
 iw,ih=Image.open(path).size
 ratio=min(w/iw,h/ih);pw,ph=iw*ratio,ih*ratio
 s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def pill(s,x,y,w,label,dark=False):
 rect(s,x,y,w,0.36,GOLD,radius=True);txt(s,x+0.13,y+0.065,w-0.26,0.25,label,11,NAVY,True)

def base(title,section,subtitle='',dark=False,source=''):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY if dark else BG)
 n=len(P.slides);color=WHITE if dark else INK
 rect(s,0,0,0.16,7.5,GOLD)
 txt(s,.58,.30,10,.25,section.upper(),11,GOLD if dark else TEAL,True)
 txt(s,.58,.77,12.15,.68,title,29,color,True)
 if subtitle:txt(s,.60,1.48,12.05,.5,subtitle,15,'CBD8E3' if dark else GRAY)
 rect(s,.60,7.01,12.1,.012,'4B6274' if dark else 'CCD7DF')
 txt(s,.60,7.13,8.8,.22,'UTFPR APucarana  •  Sistemas Embarcados  •  Pedro Reis'.replace('APucarana','Apucarana'),10,'CBD8E3' if dark else GRAY)
 txt(s,11.45,7.11,1.23,.26,f'{n:02d}  /  16',11,'CBD8E3' if dark else GRAY,align=PP_ALIGN.RIGHT)
 if source:txt(s,.60,6.73,12.0,.20,source,9.5,'CBD8E3' if dark else GRAY)
 return s

def card(s,x,y,w,h,label,body,dark=False,accent=TEAL):
 rect(s,x,y,w,h,DARKCARD if dark else WHITE,radius=True)
 rect(s,x,y,.065,h,GOLD if dark else accent)
 txt(s,x+.22,y+.18,w-.44,.55,label,20,WHITE if dark else INK,True)
 txt(s,x+.22,y+.91,w-.44,h-1.06,body,18,'DCE6ED' if dark else GRAY)

def note(s,title,seconds,text,question):
 n=len(P.slides)
 full=f'{title}\n\nTempo sugerido: {seconds}\n\n{text}\n\nPonto para a arguição: {question}'
 s.notes_slide.notes_text_frame.text=full
 SLIDES.append({'numero':n,'titulo':title,'tempo':seconds,'fala':text,'arguicao':question})

# 1 — Capa
s=base('Relógio e alarme com FreeRTOS','Exame de Suficiência · 2026.2',dark=True)
pill(s,.62,1.70,3.42,'ARDUINO MEGA 2560  ·  LCD 16×2')
txt(s,.62,2.38,7.70,1.18,'Quatro tarefas.\nUm único dono do estado.',31,WHITE,True)
txt(s,.65,3.92,7.10,.85,'Operação local, comunicação por filas\ne contagem de tempo em Timer1.',22,'CBD8E3')
txt(s,.65,5.49,7.8,.89,'Pedro Lucas dos Reis Silva  ·  RA 2272040\nEngenharia de Computação · UTFPR Apucarana\nProfessor Adalberto Lazarini',15,'CBD8E3')
photo(s,'02_alarme_habilitado.jpg',9.05,1.45,3.15,5.21)
note(s,'Relógio e alarme com FreeRTOS','30 s','Apresentar o produto como um relógio autônomo configurado pelo joystick. Explicar que o computador fornece somente alimentação durante a operação. Anunciar as três decisões que orientam a apresentação: separar tarefas, concentrar o estado no controlador e usar Timer1 para a contagem do calendário. A foto mostra o protótipo utilizado na apresentação.','Não confundir operação autônoma com referência de tempo de alta exatidão.')

# 2 — Requisitos e produto
s=base('O protótipo atende aos quatro requisitos','01 · Produto','A configuração e a operação acontecem no próprio dispositivo.',source='Fonte: Suficiencia_2026_2.pdf · fotografias do autor, 08/10/2026')
card(s,.62,2.10,4.05,1.78,'01  Exibição local','Data e hora no LCD 16×2.')
card(s,4.91,2.10,4.05,1.78,'02  Ajuste por joystick','Hora, segundos, data e alarme.')
card(s,.62,4.10,4.05,1.92,'03  Alarme','Mensagem no LCD, buzzer e LED.')
card(s,4.91,4.10,4.05,1.92,'04  Soneca','Duração configurável: 1 a 10 min.')
photo(s,'02_alarme_habilitado.jpg',9.55,1.92,2.48,4.72)
note(s,'Requisitos e produto','40 s','Relacionar cada função ao enunciado do professor. Mostrar no protótipo o LCD, o joystick e o botão D5. A alimentação USB não fornece data e hora ao programa. O potenciômetro pertence ao contraste do LCD, não ao volume. O firmware permite ajustar segundos para acertar o horário pela interface. A operação não exige terminal nem aplicação no computador.','A0 e A1 são constantes numéricas equivalentes a 54 e 55 no Mega; não são strings.')

# 3 — RTOS vs bare metal
s=base('Por que usar RTOS neste relógio?','02 · Fundamentos','O RTOS organiza a concorrência; a máquina de estados organiza o comportamento.')
card(s,.62,2.12,5.82,2.88,'Laço principal / bare metal','A aplicação define a sequência.\nMenor custo de memória.\nTambém funciona com estados e ISRs.')
card(s,6.70,2.12,5.98,2.88,'FreeRTOS preemptivo','Prioridades e esperas pelo kernel.\nFilas para trocar mensagens.\nCusto: pilhas e troca de contexto.')
rect(s,.62,5.32,12.06,.91,PALE,radius=True)
txt(s,.90,5.55,11.47,.46,'Um núcleo: as tarefas alternam a CPU; não executam em paralelo.',22,TEAL,True)
note(s,'RTOS, bare metal e máquina de estados','55 s','Explicar que uma solução bare metal bem projetada também poderia atender ao problema. O RTOS foi solicitado no exame e permite atribuir responsabilidades e prioridades explícitas. Quando o alarme fica pronto, ele pode interromper a apresentação no LCD. A máquina de estados continua presente dentro do controlador para decidir o significado de cada evento. O RTOS não aumenta a frequência da CPU nem melhora sozinho a precisão do oscilador.','Preempção é a retirada da CPU de uma tarefa pronta de menor prioridade quando outra de maior prioridade deve executar.')

# 4 — Arquitetura
s=base('O estado tem um único dono','03 · Arquitetura','As tarefas comunicam intenção e snapshots, sem compartilhar o calendário para escrita.',True,source='Fonte: RelogioFinal.ino · tarefaRelogio(), publicar(), enviar()')
rect(s,4.55,2.05,3.55,.72,DARKCARD,radius=True);txt(s,4.72,2.22,3.20,.35,'ISR Timer1 → contador',18,WHITE,True,align=PP_ALIGN.CENTER)
arrow(s,6.11,2.87,.33,.48,GOLD,True)
rect(s,.65,3.51,2.58,1.20,DARKCARD,radius=True);txt(s,.87,3.75,2.12,.8,'Botões\ne joystick',22,WHITE,True,align=PP_ALIGN.CENTER)
arrow(s,3.49,3.98,.68,.24,GOLD);txt(s,3.32,3.29,1.05,.62,'FIFO\n16 eventos',11,'CBD8E3',align=PP_ALIGN.CENTER)
rect(s,4.53,3.36,3.65,1.80,TEAL,radius=True);txt(s,4.75,3.67,3.18,.46,'TaskRelógio',26,WHITE,True,align=PP_ALIGN.CENTER);txt(s,4.75,4.32,3.18,.38,'Estado privado',19,WHITE,align=PP_ALIGN.CENTER)
arrow(s,8.47,3.64,.65,.23,GOLD);arrow(s,8.47,4.60,.65,.23,GOLD)
rect(s,9.38,3.24,3.0,.94,DARKCARD,radius=True);txt(s,9.60,3.44,2.55,.52,'Alarme  ·  ligar/desligar',17,WHITE,True)
rect(s,9.38,4.40,3.0,.94,DARKCARD,radius=True);txt(s,9.60,4.61,2.55,.52,'Display  ·  snapshot',18,WHITE,True)
txt(s,9.25,5.57,3.25,.43,'Saídas: caixas de 1 posição',13,'CBD8E3',align=PP_ALIGN.CENTER)
txt(s,.80,6.13,11.75,.35,'Quem lê o botão não decide a regra. Quem escreve no LCD não altera a hora.',20,WHITE,align=PP_ALIGN.CENTER)
note(s,'Arquitetura por eventos','60 s','Percorrer as setas da esquerda para a direita. A tarefa de botões produz um evento, como Selecionar ou Soneca. Somente a tarefa de relógio decide como esse evento muda o estado. Ela envia ao alarme um booleano e ao display uma cópia consistente do estado. A ISR incrementa um contador; o controlador o lê atomicamente. O termo dono único se refere ao estado civil, não à ausência de toda variável global no programa. Os handles das filas e o contador da ISR continuam existindo.','Copiar uma estrutura pela fila só isola os dados se ela não depender de ponteiros para objetos mutáveis externos; o snapshot deste projeto contém dados por valor.')

# 5 — Prioridades
s=base('Prioridade define quem executa quando há disputa','04 · Tarefas','A tarefa de maior prioridade também precisa liberar a CPU.')
cols=[.76,3.66,5.13,9.42]
for x,w,label in [(cols[0],2.7,'TAREFA'),(cols[1],1.35,'PRIOR.'),(cols[2],4.1,'RESPONSABILIDADE'),(cols[3],2.8,'COMO AGUARDA')]:txt(s,x,2.09,w,.32,label,12,TEAL,True)
rows=[('Alarme','3','Buzzer e LED','Fila / fase sonora'),('Relógio','2','Calendário e regras','vTaskDelayUntil'),('Botões','2','Leitura e eventos','vTaskDelayUntil'),('Display','1','LCD e diagnóstico','Fila de snapshots')]
for i,row in enumerate(rows):
 y=2.62+i*.77;rect(s,.62,y-.09,12.06,.65,WHITE if i%2==0 else LIGHT,radius=True)
 for j,(x,val,w) in enumerate(zip(cols,row,[2.7,1.35,4.1,2.8])):txt(s,x,y+.03,w,.43,val,19,INK,j==0)
txt(s,.80,6.04,11.9,.49,'Bloqueada no kernel ≠ esperando em um laço que ocupa a CPU.',22,TEAL,True)
note(s,'Tarefas e prioridades','55 s','Justificar a prioridade 3 do alarme pela resposta à ativação e ao desligamento. Explicar que ele fica bloqueado quando inativo e aguarda a próxima fase quando está tocando. Relógio e botões têm prioridade 2; o port habilita divisão de tempo entre tarefas prontas de mesma prioridade. O display tem prioridade 1 e pode ser preemptado. A prioridade não é uma porcentagem reservada de CPU. A espera em fila entrega a CPU a outra tarefa; ela não é busy waiting.','O display não retém um mutex de aplicação que a tarefa de relógio precise adquirir.')

# 6 — Filas/mutex
s=base('Mutex protege acesso; fila transporta informação','05 · Comunicação','A evolução central foi retirar a escrita compartilhada do estado civil.',source='Fontes: versões preservadas · queue.h / queue.c · documentação de mutexes do FreeRTOS')
card(s,.62,2.10,5.82,2.57,'Antes: mutex','Adquirir → acessar → liberar.\nEstrutura compartilhada.\nHerança de prioridade no kernel.')
card(s,6.70,2.10,5.98,2.57,'Agora: filas','Evento → controlador → cópia.\nEstado privado do relógio.\nLCD sem lock compartilhado.')
rect(s,.62,4.98,5.82,1.22,PALE,radius=True);txt(s,.89,5.15,5.25,.85,'Entrada: FIFO de 16 eventos\nPreserva a ordem dos comandos aceitos.',18,TEAL)
rect(s,6.70,4.98,5.98,1.22,PALE,radius=True);txt(s,6.96,5.15,5.42,.85,'Saídas: xQueueOverwrite\nConserva o estado mais recente.',18,TEAL)
note(s,'Mutex versus filas','65 s','Explicar que o mutex é adequado quando há um recurso realmente compartilhado. O problema não é a existência do mutex, mas a duração e a disciplina da exclusividade. A herança de prioridade reduz a interferência de tarefas intermediárias quando alguém de maior prioridade espera pelo recurso. Nesta arquitetura, o display recebe uma cópia e não segura um lock necessário ao controle. A fila de entrada preserva ordem; as caixas de saída substituem dados antigos porque o consumidor precisa do estado atual.','Uma fila não garante ausência de toda condição de corrida. A segurança aqui depende de dono único, cópia de dados, atomicidade na ISR e política explícita para fila cheia.')

# 7 — Timers
s=base('Tick do RTOS e relógio civil são bases distintas','06 · Temporização','Timer0, Timer1 e Timer2 têm funções diferentes nesta implementação.',True,source='Fontes: iniciarTempo() · FreeRTOSVariant.h · core AVR: wiring.c e Tone.cpp')
card(s,.64,2.16,5.85,1.79,'Watchdog → escalonador','Tick nominal ≈ 16 ms.\nAgenda esperas e tarefas.',True)
card(s,6.76,2.16,5.86,1.79,'Timer1 → calendário','CTC a 100 Hz.\nUma contagem nominal a cada 10 ms.',True)
card(s,.64,4.20,5.85,1.79,'Timer0 → millis()','Base do core Arduino.\nDiagnóstico; não é o RTC do projeto.',True)
card(s,6.76,4.20,5.86,1.79,'Timer2 → tone()','Geração da onda de áudio.\nBuzzer a 2 kHz durante a fase ativa.',True)
note(s,'Domínios temporais','65 s','Distinguir o timer do escalonador da referência usada para avançar o calendário. O watchdog permanece no tick do FreeRTOS. Timer1 utiliza o clock principal da placa em CTC, a 100 Hz nominal. Timer0 mantém millis e micros do Arduino e aparece no diagnóstico. Timer2 é usado por tone para o buzzer. O firmware não usa Timer0 como um RTC de calendário. Trocar para uma biblioteca de RTC no Timer2 exigiria tratar o conflito com tone.','pdMS_TO_TICKS(20) produz um tick nesta configuração de 62 Hz; isso não transforma o período físico em exatamente 20 ms.')

# 8 — ISR atomicidade
s=base('A ISR registra tempo; a tarefa atualiza o calendário','07 · Concorrência com interrupções','Prescaler = 64  ·  OCR1A = 2499  ·  clock nominal = 16 MHz',source='Trechos de RelogioFinal.ino · lerTempo() e tarefaRelogio()')
rect(s,.64,2.13,12.01,.73,PALE,radius=True);txt(s,.9,2.29,11.47,.42,'16.000.000 / [64 × (2499 + 1)] = 100 Hz',25,TEAL,True,align=PP_ALIGN.CENTER)
rect(s,.64,3.15,5.84,2.45,NAVY,radius=True)
txt(s,.89,3.37,5.38,2.10,'ISR(TIMER1_COMPA_vect) {\n  ++ticks100Hz;\n}\n\nATOMIC_BLOCK(ATOMIC_RESTORESTATE) {\n  copia = ticks100Hz;\n}',15,WHITE,font=MONO)
card(s,6.76,3.15,5.89,2.45,'Leitura íntegra no AVR','Contador de 32 bits em CPU de 8 bits.\nvolatile não basta: leitura atômica.\nLCD fora da seção crítica.')
rect(s,.64,5.88,12.01,.59,LIGHT,radius=True);txt(s,.90,6.00,11.44,.36,'agora − anterior  →  contagens decorridas  →  calendário e prazos',21,INK,True,align=PP_ALIGN.CENTER)
note(s,'Timer1 e atomicidade','65 s','Apresentar a fórmula do CTC e apontar que todos os valores são nominais. A ISR faz apenas o incremento, mantendo baixo o trabalho em interrupção. No AVR, a leitura de quatro bytes pode ser interrompida entre bytes; volatile não resolve isso. ATOMIC_RESTORESTATE protege somente a cópia e restaura o estado anterior. O controlador usa a subtração unsigned para recuperar as contagens desde a amostra anterior, inclusive atravessando rollover. Um atraso no agendamento não precisa se tornar perda de segundos.','Esse mecanismo recupera contagens registradas. Não reconstrói interrupções perdidas se o hardware deixar de distingui-las durante um bloqueio excessivo.')

# 9 — Interface
s=base('Editar o alarme não interrompe o relógio','08 · Interface por eventos','A edição usa um rascunho; somente o controlador aplica a configuração.',source='Fonte: RelogioTipos.h · Botao::ler(), Eixo::ler() · comandar()')
photo(s,'04_edicao_segundos.jpg',.68,2.03,2.16,4.39)
photo(s,'08_edicao_alarme.jpg',3.00,2.03,2.16,4.39)
card(s,5.62,2.07,6.97,2.29,'Navegação','← / →  escolhe o campo\n↑ / ↓  altera o valor\nSEL salva  ·  D5 cancela a edição')
rect(s,5.62,4.63,6.97,1.76,PALE,radius=True)
txt(s,5.89,4.82,6.40,1.41,'Debounce: 50 ms de estabilidade.\nSEL longo: um evento após 1 s.\nSem esperar o botão ser solto em um laço.',19,TEAL)
note(s,'Entrada e edição','55 s','Mostrar a foto da edição de segundos e a do horário do alarme. A tarefa de botões filtra a entrada e gera eventos sem decidir a regra do menu. O rascunho fica separado do relógio corrente. Salvar apenas o alarme mantém a hora que continuou avançando; salvar uma alteração da hora aplica o valor escolhido e reinicia a fração civil. O clique curto de SEL é emitido na soltura e o longo não gera um curto adicional.','Os 50 ms de debounce são estabilidade observada na base Timer1, não apenas um delay depois de qualquer leitura.')

# 10 — Alarme soneca
s=base('A soneca usa duração relativa, não hora do dia','09 · Máquina de estados','O mesmo evento de soneca pode vir de D5 ou do clique em SEL.',source='Fonte: comandar(), consumirPrazo(), avaliar() · exemplo de meia-noite exercitado nos testes')
for x,title,small in [(.66,'Silêncio','Aguarda o horário habilitado'),(4.58,'Tocando','Toque de até 60 segundos'),(8.50,'Soneca','Espera de 1 a 10 minutos')]:
 rect(s,x,2.43,3.55,1.29,TEAL if title=='Tocando' else WHITE,radius=True)
 txt(s,x+.16,2.63,3.23,.43,title,25,WHITE if title=='Tocando' else INK,True,align=PP_ALIGN.CENTER)
 txt(s,x+.16,3.22,3.23,.35,small,12,WHITE if title=='Tocando' else GRAY,align=PP_ALIGN.CENTER)
arrow(s,4.25,2.96,.29,.19);arrow(s,8.17,2.96,.29,.19)
txt(s,3.76,2.01,1.39,.34,'hora : min : 00',11,GRAY,align=PP_ALIGN.CENTER)
txt(s,7.72,2.01,1.1,.34,'D5 / SEL',11,GRAY,align=PP_ALIGN.CENTER)
rect(s,4.58,3.92,7.47,.66,LIGHT,radius=True);txt(s,4.81,4.08,7.03,.32,'Fim da soneca → novo toque; SEL longo → encerra.',17,INK)
rect(s,.66,4.96,11.39,1.31,PALE,radius=True)
txt(s,.91,5.15,10.91,.47,'23:59:50  +  60 segundos  →  00:00:50',27,TEAL,True,align=PP_ALIGN.CENTER)
txt(s,.91,5.83,10.91,.28,'A mudança de dia não invalida o prazo restante.',16,TEAL,align=PP_ALIGN.CENTER)
note(s,'Alarme e soneca','65 s','Percorrer os estados. No segundo zero do horário habilitado, o controlador inicia o toque. D5 ou SEL curto inicia a mesma soneca configurada. A contagem é relativa, por isso a meia-noite não quebra a comparação. O exemplo de 60 segundos usa a configuração mínima de um minuto; o padrão é cinco minutos. Cada toque termina em até 60 segundos. A repetição da soneca requer novo acionamento; ela não se repete automaticamente para sempre.','D5 durante a soneca cancela e desabilita o alarme. SEL longo encerra o aviso, mas preserva a habilitação diária. A chave de ocorrência evita redisparo no mesmo minuto e data.')

# 11 — Precisão
s=base('Precisão física não é a mesma coisa que prioridade','10 · Precisão','Três fenômenos precisam ser separados na análise do horário.',True)
card(s,.66,2.26,3.83,2.75,'Acerto do horário','Diferença de horário\nno instante em que\no relógio é acertado.',True)
card(s,4.76,2.26,3.83,2.75,'Deriva','Erro que cresce\ncom a frequência real\ndo oscilador.',True)
card(s,8.86,2.26,3.81,2.75,'Latência','Tempo até a tarefa\nprocessar um evento\nou atualizar a saída.',True)
txt(s,.85,5.48,11.64,.98,'Timer1 separa o calendário do watchdog.\nA exatidão continua dependendo da referência física.',24,WHITE,True,align=PP_ALIGN.CENTER)
note(s,'Precisão e exatidão','60 s','Explicar que prioridade organiza resposta da CPU; ela não calibra o oscilador. A diferença observada na aferição para apresentação não determina sozinha a deriva. Para medir deriva, compara-se a mudança do erro em um intervalo conhecido. Timer1 usa o oscilador principal e pode adiantar ou atrasar. Uma melhoria de referência poderia usar RTC externo ou Timer2 assíncrono com cristal apropriado, mas não faz parte desta montagem. Não apresentar essa alternativa como algo implementado.','Timer2 assíncrono não é a única solução e também tem tolerância. A montagem atual não dispõe de sincronização automática com uma fonte independente de horário.')

# 12 — Resultados
s=base('O comportamento foi verificado em etapas','11 · Resultados','Montagem funcional, núcleo testado e firmware compilado para o Mega.',source='Fontes: testes/nucleo_teste.cpp · compilação Arduino AVR 1.8.8 · fotografias do protótipo')
rect(s,.66,2.20,3.25,2.31,TEAL,radius=True);txt(s,.96,2.54,2.65,.94,'13',60,WHITE,True,align=PP_ALIGN.CENTER);txt(s,.92,3.62,2.73,.60,'grupos de testes\ndo núcleo aprovados',18,WHITE,align=PP_ALIGN.CENTER)
card(s,4.20,2.20,4.03,2.31,'Calendário e prazos','Meia-noite e ano bissexto.\nSoneca e fim do toque.\nRollover e recuperação.')
card(s,8.49,2.20,4.12,2.31,'Entrada e display','Edição de segundos.\nDebounce e clique longo.\nLimites do LCD 16×2.')
rect(s,.66,4.84,11.95,1.32,WHITE,radius=True)
txt(s,.95,5.04,11.37,.52,'18.388 bytes de programa  ·  729 bytes de SRAM estática',23,INK,True,align=PP_ALIGN.CENTER)
txt(s,.95,5.77,11.37,.28,'Pilhas e filas acrescentam alocação em execução; o ensaio temporal do núcleo usa contagens nominais.',12,GRAY,align=PP_ALIGN.CENTER)
note(s,'Resultados e verificação','55 s','Apresentar a evidência adequada para cada conclusão. As fotos mostram a montagem e a interface; a compilação confirma a construção para ATmega2560; os testes verificam a lógica extraída do próprio sketch. O teste de seis horas trabalha com 2.160.000 contagens nominais, não é uma certificação de deriva física de seis horas. O consumo estático de 729 bytes não inclui todas as alocações do kernel. A integração mantém verificações de criação de recursos e diagnóstico de margem de pilha.','Não transformar o número de bytes estáticos em uma afirmação de memória total livre; há pilhas, filas, TCBs e outras estruturas em execução.')

# 13 — Conclusão / demonstração
s=base('Demonstração: configurar, disparar e adiar','12 · Encerramento','O projeto integra tarefas, eventos e interrupções em uma interface autônoma.',True)
for y,num,title,desc in [(2.20,'1','Configurar pelo joystick','Hora, segundos, alarme e soneca.'),(3.56,'2','Observar o disparo','Mensagem no LCD, LED e buzzer.'),(4.92,'3','Acionar D5 e encerrar','Soneca no LCD; SEL longo encerra o aviso.')]:
 pill(s,.68,y+.1,.55,num);txt(s,1.48,y,7.37,.48,title,25,WHITE,True);txt(s,1.49,y+.65,7.37,.46,desc,18,'CBD8E3')
photo(s,'09_edicao_soneca.jpg',9.15,2.03,2.57,4.49)
note(s,'Conclusão e demonstração','40 s + demonstração','Concluir com os três pontos centrais: um dono para o estado, comunicação por mensagens e separação das bases de tempo. Na bancada, abrir a edição, navegar até o horário do alarme, configurá-lo para o próximo minuto e salvar. Habilitar pelo D5 na tela normal. Enquanto aguarda, mostrar que o relógio continua avançando. No disparo, pressionar D5 para exibir a soneca. Encerrar com SEL longo, ou explicar que D5 durante a soneca também desabilita o alarme. Não reinicializar a placa durante a demonstração, pois a configuração é volátil.','Para uma demonstração curta, configurar soneca de um minuto. A exposição não precisa esperar todo esse minuto: pode mostrar a contagem e encerrar o aviso.')

# 14 — Apoio: registradores
s=base('Timer1: valores e significado dos registradores','Apoio A · Temporização','CTC a 100 Hz · divisão do clock principal de 16 MHz.',source='Fonte: iniciarTempo() · RelogioFinal.ino')
rows=[('TCCR1A = 0','Desconecta as saídas OC1 do controle do timer.'),('TCCR1B: WGM12','Seleciona CTC.'),('TCCR1B: CS11 | CS10','Seleciona prescaler 64.'),('OCR1A = 2499','Comparação nominal a 100 Hz.'),('TIFR1: OCF1A = 1','Limpa uma comparação pendente.'),('TIMSK1: OCIE1A = 1','Habilita a interrupção de comparação A.')]
for i,(a,b) in enumerate(rows):
 y=2.08+i*.58;rect(s,.66,y,12.00,.53,WHITE if i%2==0 else LIGHT)
 txt(s,.88,y+.11,4.17,.35,a,16,INK,True,font=MONO);txt(s,5.39,y+.10,7.0,.36,b,17,GRAY)
txt(s,.85,6.02,11.75,.53,'Trocar a divisão para 1 Hz não corrige o erro do mesmo oscilador.',21,TEAL,True)
note(s,'Apoio: registradores','Consultar se solicitado','Explicar a ordem de inicialização: parar o clock, desabilitar interrupção do periférico, limpar controle e contador, configurar comparação, limpar flag, habilitar interrupção e ligar o timer. A sequência ocorre em bloco atômico antes das tarefas. Para prescaler 1024 e OCR1A 15624, a divisão nominal seria 1 Hz, mas a unidade de todos os contadores teria de ser revista. Como a fonte de clock continuaria a mesma, a tolerância física também permaneceria.','F_CPU é uma constante de compilação, não uma medição automática da frequência real da placa.')

# 15 — Apoio: perguntas
s=base('Quatro respostas importantes para a arguição','Apoio B · Concorrência e memória','As garantias pertencem à arquitetura completa, não a uma API isolada.')
card(s,.65,2.10,5.88,1.92,'“Fila elimina qualquer corrida?”','Não. Dono único + cópia + leitura\natômica sustentam a proteção.')
card(s,6.78,2.10,5.88,1.92,'“Nunca há bloqueio?”','Há espera pelo kernel.\nNão há espera ocupada nos botões.')
card(s,.65,4.28,5.88,1.92,'“E se a entrada ficar cheia?”','Envio sem espera; perda notificada\ne mensagem para repetir o comando.')
card(s,6.78,4.28,5.88,1.92,'“E a memória das tarefas?”','1.568 bytes nas quatro pilhas.\nMargem observada por high-water mark.')
note(s,'Apoio: perguntas da banca','Consultar se solicitado','Esclarecer que ausência de mutex compartilhado remove aquele caminho de inversão de prioridade, sem prometer latência nula ou impossibilidade de qualquer defeito. As filas possuem proteção interna e capacidade finita. O snapshot de saída pode substituir outro porque interessa o estado atual. A notificação contadora registra evento de entrada descartado. O high-water mark informa a menor margem de pilha observada no percurso executado, não uma prova de todos os percursos possíveis.','StackType_t é uint8_t no port AVR usado: profundidade 512 corresponde a 512 bytes. Alarme 256 + relógio 512 + botões 288 + display 512 = 1.568 bytes.')

# 16 — Referências
s=base('Fontes e material da entrega','Apoio C · Referências','A teoria acompanha o código da biblioteca e o firmware efetivamente entregue.')
items=[('Enunciado do exame','Adalberto Lazarini · UTFPR Apucarana · 2026.2','https://github.com/plreis/Sistemas-Embarcados/blob/main/documentos/Suficiencia_2026_2.pdf'),('Arduino FreeRTOS 11.1.0-3','tasks.c · queue.c · queue.h · FreeRTOSVariant.h','https://github.com/feilipu/Arduino_FreeRTOS_Library'),('Arduino AVR Core 1.8.8','wiring.c · Tone.cpp','https://github.com/arduino/ArduinoCore-avr'),('AVR-LibC e Microchip','ATOMIC_BLOCK e configuração de interrupções megaAVR','https://avrdudes.github.io/avr-libc/avr-libc-user-manual/group__util__atomic.html'),('Relatório, firmware e fotografias','github.com/plreis/Sistemas-Embarcados','https://github.com/plreis/Sistemas-Embarcados')]
for i,(a,b,url) in enumerate(items):
 y=2.04+i*.86
 z=txt(s,.77,y,11.85,.35,a,21,INK,True)
 z.text_frame.paragraphs[0].runs[0].hyperlink.address=url
 txt(s,.78,y+.41,11.80,.29,b,14,GRAY)
note(s,'Referências','Consultar se solicitado','O repositório reúne o relatório, o firmware consolidado, as fotos originais e as versões anteriores. A pasta referencias preserva os excertos reais do FreeRTOS e a origem dos trechos. A apresentação usa o mesmo conteúdo técnico da entrega e mantém a distinção entre comportamento do software, observação do produto e precisão da referência física.','O relatório possui as referências completas e a comparação entre a versão com mutex e a versão por eventos.')

P.save(BASE/'Apresentacao_Relogio_RTOS.pptx')
route=['# Roteiro de apresentação — Relógio e alarme com FreeRTOS','',
'13 slides principais para cerca de 12 minutos de exposição, além da demonstração. Os slides 14–16 são apoio para a arguição. O mesmo roteiro está nas notas do apresentador do PowerPoint.','',
'## Sequência de fala','']
for item in SLIDES:
 route += [f'### {item["numero"]:02d}. {item["titulo"]}',f'**Tempo sugerido:** {item["tempo"]}.','',item['fala'],'', '**Ponto para a arguição:** '+item['arguicao'],'']
route += ['## Preparação da demonstração','',
'- Manter a placa ligada: a configuração permanece em SRAM e volta ao estado inicial após reinicialização.',
'- Confirmar que a hora exibida está acertada e o contraste permite leitura pela banca.',
'- Configurar o alarme para o próximo minuto e salvar; habilitar por D5 na tela normal.',
'- Durante o toque, D5 ou SEL curto inicia a soneca. SEL longo encerra o aviso.',
'- Durante a soneca, D5 cancela e também desabilita o alarme; explicar essa diferença.',
'- Usar o computador somente como alimentação durante a demonstração. O monitor serial não integra a operação.',
'- Manter a apresentação e o PDF disponíveis localmente.','',
'## Respostas que exigem precisão','',
'**Por que não houve sincronização automática?** A montagem não dispõe de uma fonte independente de horário acessível ao firmware. Timer1 mede intervalos do próprio oscilador; não determina sozinho se a hora está correta. Uma nova referência exigiria integração adicional.',
'', '**Por que vTaskDelayUntil e não apenas vTaskDelay?** O primeiro mantém a referência periódica a partir do despertar planejado anterior; o segundo espera relativamente ao instante da chamada. O calendário desta aplicação depende do contador Timer1, não da quantidade de despertares.',
'', '**Como funciona o rollover?** A diferença unsigned calcula as contagens entre amostras, desde que não transcorra um ciclo completo do contador entre elas. A 100 Hz, 32 bits dão aproximadamente 497 dias. Esse contador é distinto do tick de 16 bits do kernel.',
'', '**Pode chamar o sistema de imune à inversão de prioridade?** A formulação precisa é: o display não retém mutex de aplicação necessário ao relógio ou ao alarme. Isso remove esse caminho de inversão. Ainda existem escalonamento, seções internas do kernel e latência de interrupção.',
'', '**O que a aferição para apresentação demonstra?** Uma comparação pontual do horário. A taxa de deriva exige comparar a mudança do erro durante um intervalo conhecido.',
'', '**A soneca repete para sempre?** Não há repetição automática infinita. Cada novo adiamento depende de D5 ou SEL enquanto o alarme toca. O toque tem duração máxima de 60 segundos.',
'', '**Existe proteção contra falta de memória?** A criação de cada recurso tem o retorno verificado. Os recursos são criados na inicialização, sem alocação repetitiva nos ciclos normais. A biblioteca habilita verificação de overflow e o diagnóstico acompanha a margem de pilha.','']
(BASE/'Roteiro_apresentacao.md').write_text('\n'.join(route))
print(f'Gerados {len(P.slides)} slides com notas do apresentador.')
